"""Build outreach pools for the final list (final_5y_under_200m.csv).

Assignments come from ../pools/assignments.psv (earlier hand-made, reused for companies still in the list) and
todo_batch*_assignments.psv (new). Voice AI companies go to VOICE_EXCLUDED and are not in any outreach pool.
Outputs: pools.csv, pool_definitions.csv, voice_excluded.csv, pools.xlsx, validation printout.
"""
import csv
import glob
import os
import re
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
LEAD = "We are Fast Code AI, an applied ML lab, and we are currently working with startups and enterprises on "
TAIL = ", which is very close to what {company} needs."
POOLS = [  # pool_id, sheet name, pool_name, proof_points, sentence middle, nameable clients
    ("PHYS", "P2 Physical AI", "Perception and physical AI", "P2",
     "perception and vision-language models for vehicles and robots, including systems that run in production cars today",
     "Bosch; Mercedes-Benz (MBUX); CausalDriveBench (NeurIPS)"),
    ("PHYS_EVAL", "P2+P3 Robot data & evals", "Physical AI data and evaluation", "P2 + P3",
     "perception models for physical AI and the evaluation systems that measure how well they work in the real world",
     "Bosch; Mercedes-Benz; CausalDriveBench (NeurIPS); ThoughtSpot; Entelligence; Tattvam AI"),
    ("AGENT_EVAL", "P3 Agent quality", "Agent quality and evaluation", "P3",
     "evaluation harnesses and quality systems that keep AI agents reliable in production",
     "ThoughtSpot; Entelligence; Tattvam AI"),
    ("CODE", "P3 Coding agents", "Coding agents and code review (P3 coding facet)", "P3",
     "coding agents and automated pull request review for engineering teams",
     "Entelligence; ThoughtSpot; Tattvam AI"),
    ("RAG", "P4 Retrieval", "Retrieval over large document sets", "P4",
     "retrieval systems that find the right answer across very large enterprise document collections",
     "ThoughtSpot; MIAI"),
    ("LEGAL", "P4+P5 Legal AI", "Legal AI over legal documents", "P4 + P5",
     "legal AI systems built with a law firm, including retrieval across large sets of legal documents",
     "MIAI; ThoughtSpot"),
    ("SEARCH", "P6 Search & optimization", "Search and optimization agents", "P6",
     "LLM agents that search for better designs by optimizing against slow and expensive simulators",
     "Tattvam AI"),
    ("DIAGRAM", "P7 Industrial diagrams", "Engineering and industrial diagrams", "P7",
     "AI systems that digitize engineering and process diagrams for large industrial companies",
     "Saudi Aramco"),
    ("GENERIC", "Generic", "Generic (no proof point)", "",
     "hard ML research and engineering problems, embedding senior ML engineers directly in their teams",
     "None"),
]
PDEF = {p[0]: p for p in POOLS}
VALID = set(PDEF) | {"VOICE_EXCLUDED"}
COLS = ["company_key", "company", "batch", "website", "pool_id", "pool_name", "proof_points", "secondary_proof",
        "fit_strength", "core_problem", "reason", "runner_up_pool", "opener_hint", "proof_sentence",
        "latest_round_date", "latest_round_amount_usd", "total_funding_usd", "funding_source", "open_positions",
        "ml_jobs", "team_size", "one_liner"]


def load_psv(path, rename_voice=False):
    out = {}
    for line in open(path):
        if not line.strip():
            continue
        p = line.rstrip("\n").split("|")
        if len(p) != 8:
            print("BAD LINE", path, p[:2])
            continue
        if rename_voice and p[1] == "VOICE":
            p[1] = "VOICE_EXCLUDED"
        out[p[0]] = p
    return out


def main():
    comp = {r["key"]: r for r in csv.DictReader(open(os.path.join(HERE, "companies_for_pools.csv")))}
    final = {r["website"]: r for r in csv.DictReader(open(os.path.join(ROOT, "final_5y_under_200m.csv")))}
    prev = load_psv(os.path.join(ROOT, "pools", "assignments.psv"), rename_voice=True)
    voice_prev = set(open(os.path.join(ROOT, "pools", "excluded_voice.txt")).read().split())
    assign = {k: v for k, v in prev.items() if k in comp}
    for k in voice_prev & set(comp):
        assign[k][1], assign[k][3], assign[k][2] = "VOICE_EXCLUDED", "None", ""
    for p in sorted(glob.glob(os.path.join(HERE, "todo_batch*_assignments.psv"))):
        for k, v in load_psv(p).items():
            if k in comp and k not in prev:
                assign[k] = v
    # Validation
    problems = []
    missing = sorted(set(comp) - set(assign))
    if missing:
        problems.append(f"{len(missing)} companies without an assignment: {missing[:10]}")
    for k, v in assign.items():
        if v[1] not in VALID:
            problems.append(f"{k}: invalid pool {v[1]}")
        if "—" in "|".join(v):
            problems.append(f"{k}: em dash")
        if v[1] in ("GENERIC", "VOICE_EXCLUDED") and v[3] != "None":
            v[3] = "None"
        if v[1] not in ("GENERIC", "VOICE_EXCLUDED") and v[3] not in ("Strong", "Moderate"):
            problems.append(f"{k}: fit {v[3]}")
        if len(v[7].split()) > 26:
            problems.append(f"{k}: long opener")
    rows, voice = [], []
    for k, v in assign.items():
        c = comp[k]
        f = final.get(c["website"], {})
        pid = v[1]
        r = {"company_key": k, "company": c["company"], "batch": c["batch"], "website": c["website"], "pool_id": pid,
             "pool_name": PDEF[pid][2] if pid in PDEF else "Voice AI (excluded from outreach)",
             "proof_points": PDEF[pid][3] if pid in PDEF else "", "secondary_proof": v[2].replace("P1", "").strip(" ,+"),
             "fit_strength": v[3], "core_problem": v[5], "reason": v[6], "runner_up_pool": v[4], "opener_hint": v[7],
             "proof_sentence": (LEAD + PDEF[pid][4] + TAIL).replace("{company}", c["company"]) if pid in PDEF else "",
             "latest_round_date": c["latest_round_date"], "latest_round_amount_usd": f.get("latest_round_amount_usd", ""),
             "total_funding_usd": c["total_funding_usd"], "funding_source": f.get("source", ""),
             "open_positions": f.get("open_positions", ""), "ml_jobs": c["ml_jobs"], "team_size": f.get("team_size", ""),
             "one_liner": c["one_liner"]}
        (voice if pid == "VOICE_EXCLUDED" else rows).append(r)
    order = {"Strong": 0, "Moderate": 1, "None": 2}
    ids = [p[0] for p in POOLS]
    rows.sort(key=lambda r: (ids.index(r["pool_id"]), order.get(r["fit_strength"], 3), -(float(r["total_funding_usd"] or 0))))
    cnt = Counter(r["pool_id"] for r in rows)
    for pid in ids:
        if cnt[pid] < 3:
            problems.append(f"pool {pid} has {cnt[pid]} companies (<3)")
    with open(os.path.join(HERE, "pools.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)
    with open(os.path.join(HERE, "voice_excluded.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(voice)
    defs = []
    for pid, sheet, name, pp, mid, clients in POOLS:
        m = [r for r in rows if r["pool_id"] == pid]
        defs.append({"pool_id": pid, "pool_name": name, "proof_points": pp, "proof_sentence": LEAD + mid + TAIL,
                     "nameable_clients_for_followup": clients, "company_count": len(m),
                     "strong_count": sum(r["fit_strength"] == "Strong" for r in m),
                     "moderate_count": sum(r["fit_strength"] == "Moderate" for r in m)})
    with open(os.path.join(HERE, "pool_definitions.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(defs[0]))
        w.writeheader()
        w.writerows(defs)
    write_xlsx(rows, voice, defs)
    print(f"{len(rows)} in pools, {len(voice)} voice excluded, {len(missing)} missing")
    for d in sorted(defs, key=lambda d: -d["company_count"]):
        print(f"  {d['pool_id']:11} {d['company_count']:5}  strong {d['strong_count']:4}  moderate {d['moderate_count']:4}")
    print("problems:", len(problems))
    for p in problems[:30]:
        print("  ", p)


def write_xlsx(rows, voice, defs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    F = "Arial"
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    ov.append(["Pool ID", "Pool name", "Proof points", "Proof sentence", "Nameable clients for follow-up",
               "Companies", "Strong", "Moderate", "Tab"])
    sheet = {p[0]: p[1] for p in POOLS}
    for d in sorted(defs, key=lambda d: -d["company_count"]):
        ov.append([d["pool_id"], d["pool_name"], d["proof_points"] or "None", d["proof_sentence"],
                   d["nameable_clients_for_followup"], d["company_count"], d["strong_count"], d["moderate_count"],
                   sheet[d["pool_id"]]])
    tot = ov.max_row + 1
    ov.append(["Total", "", "", "", "", sum(d["company_count"] for d in defs), sum(d["strong_count"] for d in defs),
               sum(d["moderate_count"] for d in defs), ""])
    ov.cell(tot + 2, 1, f"{len(voice)} voice AI companies are excluded from outreach (tab 'Voice excluded'). "
                        "{company} in a proof sentence is replaced with the company name. Counts are fixed values "
                        "taken from the pool tabs when the file was built (pools_v2/build_pools_v2.py).")
    for col, wdt in zip("ABCDEFGHI", [12, 34, 12, 70, 40, 11, 9, 10, 24]):
        ov.column_dimensions[col].width = wdt
    widths = [18, 22, 12, 26, 11, 26, 10, 9, 10, 40, 44, 13, 60, 60, 12, 14, 14, 13, 10, 8, 9, 40]
    for name, data in [(p[1], [r for r in rows if r["pool_id"] == p[0]]) for p in POOLS] + [("Voice excluded", voice)]:
        ws = wb.create_sheet(name)
        ws.append(COLS)
        for r in data:
            vals = [r[c] for c in COLS]
            for i, c in enumerate(COLS):
                if c in ("latest_round_amount_usd", "total_funding_usd") and vals[i] not in ("", None):
                    try:
                        vals[i] = float(vals[i])
                    except ValueError:
                        pass
            ws.append(vals)
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        for row in ws.iter_rows(min_row=2):
            for i, c in enumerate(COLS):
                if c in ("latest_round_amount_usd", "total_funding_usd"):
                    row[i].number_format = '$#,##0;($#,##0);-'
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=F, size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in ws[1]:
            c.font = Font(name=F, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A2E")
        ws.auto_filter.ref = ws.dimensions if ws.title != "Overview" else f"A1:I{tot}"
    for c in ov[tot]:
        c.font = Font(name=F, size=10, bold=True)
    wb.save(os.path.join(HERE, "pools.xlsx"))


if __name__ == "__main__":
    main()
