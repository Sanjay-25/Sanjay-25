"""Build pools.csv, pool_definitions.csv and pools.xlsx from pools/assignments.psv."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
LEAD = "We are Fast Code AI, an applied ML lab, and we are currently working with startups and enterprises on "
TAIL = ", which is very close to what {company} needs."
POOLS = [  # pool_id, sheet name, pool_name, proof_points, sentence middle, nameable clients
    ("PHYS", "P2 Physical AI", "Perception and physical AI", "P2",
     "perception and vision-language models for vehicles and robots, including systems that run in production cars today",
     "Bosch; Mercedes-Benz (MBUX); CausalDriveBench (NeurIPS)"),
    ("PHYS_EVAL", "P2+P3 Robot data & evals", "Physical AI data and evaluation", "P2 + P3",
     "perception models for physical AI and the evaluation systems that measure how well they work in the real world",
     "Bosch; Mercedes-Benz; CausalDriveBench (NeurIPS); ThoughtSpot; Entelligence; Tattvam AI"),
    ("AGENT_EVAL", "P3 Agent quality", "Agent quality and evaluation", "P3",
     "evaluation harnesses and quality systems that keep AI agents reliable in production",
     "ThoughtSpot; Entelligence; Tattvam AI"),
    ("CODE", "P3 Coding agents", "Coding agents and code review (P3 coding facet)", "P3",
     "coding agents and automated pull request review for engineering teams",
     "Entelligence; ThoughtSpot; Tattvam AI"),
    ("RAG", "P4 Retrieval", "Retrieval over large document sets", "P4",
     "retrieval systems that find the right answer across very large enterprise document collections",
     "ThoughtSpot; MIAI"),
    ("LEGAL", "P4+P5 Legal AI", "Legal AI over legal documents", "P4 + P5",
     "legal AI systems built with a law firm, including retrieval across large sets of legal documents",
     "MIAI; ThoughtSpot"),
    ("SEARCH", "P6 Search & optimization", "Search and optimization agents", "P6",
     "LLM agents that search for better designs by optimizing against slow and expensive simulators",
     "Tattvam AI"),
    ("DIAGRAM", "P7 Industrial diagrams", "Engineering and industrial diagrams", "P7",
     "AI systems that digitize engineering and process diagrams for large industrial companies",
     "Saudi Aramco"),
    ("GENERIC", "Generic", "Generic (no proof point)", "",
     "hard ML research and engineering problems, embedding senior ML engineers directly in their teams",
     "None"),
]
PDEF = {p[0]: p for p in POOLS}
PNAME = {p[0]: p[2] for p in POOLS}
COLS = ["company_id", "company", "batch", "primary_space", "sub_space", "pool_id", "pool_name", "proof_points",
        "secondary_proof", "fit_strength", "core_problem", "reason", "runner_up_pool", "opener_hint",
        "raised_post_yc", "post_yc_amount_usd"]

space = {r["company_id"]: r for r in csv.DictReader(open(os.path.join(ROOT, "spaces", "space_map.csv")))}
res = {r["company_id"]: r for r in csv.DictReader(open(os.path.join(ROOT, "results.csv")))}
# Voice AI companies removed from outreach at the user's request.
EXCLUDE = set(open(os.path.join(HERE, "excluded_voice.txt")).read().split())
rows = []
for line in open(os.path.join(HERE, "assignments.psv")):
    if not line.strip():
        continue
    cid, pool, sec, fit, runner, core, reason, opener = line.rstrip("\n").split("|")
    if cid in EXCLUDE:
        continue
    r = res[cid]
    rows.append({"company_id": cid, "company": r["company"], "batch": r["batch"],
                 "primary_space": space[cid]["primary_space"], "sub_space": space[cid]["sub_space"],
                 "pool_id": pool, "pool_name": PNAME[pool], "proof_points": PDEF[pool][3],
                 "secondary_proof": sec, "fit_strength": fit, "core_problem": core, "reason": reason,
                 "runner_up_pool": runner, "opener_hint": opener,
                 "raised_post_yc": r["raised_post_yc"], "post_yc_amount_usd": r["post_yc_amount_usd"]})
order = {"Strong": 0, "Moderate": 1, "None": 2}
rows.sort(key=lambda r: ([p[0] for p in POOLS].index(r["pool_id"]), order[r["fit_strength"]],
                         -float(r["post_yc_amount_usd"] or 0)))

with open(os.path.join(HERE, "pools.csv"), "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(rows)

defs = []
for pid, sheet, name, pp, mid, clients in POOLS:
    members = [r for r in rows if r["pool_id"] == pid]
    defs.append({"pool_id": pid, "pool_name": name, "proof_points": pp, "proof_sentence": LEAD + mid + TAIL,
                 "nameable_clients_for_followup": clients, "company_count": len(members),
                 "strong_count": sum(r["fit_strength"] == "Strong" for r in members),
                 "moderate_count": sum(r["fit_strength"] == "Moderate" for r in members)})
with open(os.path.join(HERE, "pool_definitions.csv"), "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(defs[0])); w.writeheader(); w.writerows(defs)

# ---- workbook
F = "Arial"
hdr_font, hdr_fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A2E")
body = Font(name=F, size=10)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()
ov = wb.active; ov.title = "Overview"
ov_cols = ["Pool ID", "Pool name", "Proof points", "Proof sentence", "Nameable clients for follow-up",
           "Companies", "Strong", "Moderate", "Tab"]
ov.append(ov_cols)
sheet_of = {p[0]: p[1] for p in POOLS}
for d in sorted(defs, key=lambda d: -d["company_count"]):
    s = sheet_of[d["pool_id"]]
    r = ov.max_row + 1
    ov.append([d["pool_id"], d["pool_name"], d["proof_points"] or "None", d["proof_sentence"],
               d["nameable_clients_for_followup"],
               d["company_count"], d["strong_count"], d["moderate_count"], s])
tot = ov.max_row + 1
ov.append(["Total", "", "", "", "", sum(d["company_count"] for d in defs), sum(d["strong_count"] for d in defs),
           sum(d["moderate_count"] for d in defs), ""])
note = tot + 2
ov.cell(note, 1, "Each company is in exactly one pool. {company} in a proof sentence is replaced with the company name; "
                 "only the opener_hint changes within a pool. Counts are fixed values taken from the pool tabs when the file was built "
                 "(source: pools/assignments.psv, manual assignment, Sep 2026); rebuild with pools/build_pools.py after edits.")
for col, wdt in zip("ABCDEFGHI", [12, 34, 12, 70, 40, 11, 9, 10, 24]):
    ov.column_dimensions[col].width = wdt
for p in POOLS:
    ws = wb.create_sheet(p[1])
    ws.append(COLS)
    for r in (x for x in rows if x["pool_id"] == p[0]):
        vals = [r[c] for c in COLS]
        vals[-1] = float(r["post_yc_amount_usd"]) if r["post_yc_amount_usd"] else None
        ws.append(vals)
    widths = [22, 22, 12, 24, 26, 11, 28, 10, 10, 10, 40, 44, 14, 60, 11, 14]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        row[-1].number_format = '$#,##0;($#,##0);-'
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for c in row:
            c.font = body; c.alignment = wrap
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
    ws.auto_filter.ref = ws.dimensions if ws.title != "Overview" else f"A1:I{tot}"
for c in ov[tot]:
    c.font = Font(name=F, size=10, bold=True)
ov.cell(note, 1).font = Font(name=F, size=9, italic=True)
ov.cell(note, 1).alignment = Alignment(wrap_text=False)
wb.save(os.path.join(HERE, "pools.xlsx"))
print("built", len(rows))
